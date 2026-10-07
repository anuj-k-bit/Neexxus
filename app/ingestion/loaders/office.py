import logfire

def parse_office(file_path: str) -> str:
    """
    Parses Office documents (.docx, .pptx) using python-docx and python-pptx.
    """
    with logfire.span("Office Document Parsing", filename=file_path):
        try:
            full_text = ""
            if file_path.lower().endswith(".docx"):
                import docx
                doc = docx.Document(file_path)
                full_text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
            elif file_path.lower().endswith(".pptx"):
                from pptx import Presentation
                prs = Presentation(file_path)
                slides_text = []
                for slide in prs.slides:
                    for shape in slide.shapes:
                        if hasattr(shape, "text") and shape.text.strip():
                            slides_text.append(shape.text.strip())
                full_text = "\n".join(slides_text)
            else:
                try:
                    from unstructured.partition.auto import partition
                    elements = partition(filename=file_path)
                    full_text = "\n".join([str(el) for el in elements])
                except ImportError:
                    pass

            return full_text
        except Exception as e:
            logfire.error(f"Office Parse Failed: {e}")
            raise e
